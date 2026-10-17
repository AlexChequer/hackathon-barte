"""Setembro e outubro na mesma pasta de saída: o que um mês passa para o outro (R18, P33).

- Atlas sem contrato em setembro: as 57 transações ficam guardadas no histórico.
- Outubro com a Atlas na planilha (contrato confirmado): fatura outubro e as 57 de setembro (R18).
- Outubro ainda sem contrato: as de setembro seguem na pendência, sem se perder.
- Alerta de antecipação usa quem estava no fechamento anterior (P33); sem histórico, a regra de antes.
- De-para de EC vem do JSON de configuração, não do código.
"""

import json
from pathlib import Path

import openpyxl
import pandas as pd
import pytest
from tests.sinteticos.outubro import gerar_outubro

from tarvo import rules
from tarvo.engines.checkout_pendencias import _sem_historico
from tarvo.loaders.orbe import ARQUIVO_PENDENTES, NOME_PENDENTES, load_orbe
from tarvo.orquestrador import run_closing

DATA = Path(__file__).resolve().parents[1] / "data" / "raw" / "set2026"
ATLAS_EC, ATLAS_CNPJ = "1076801732", "39566044000119"


def _ids_novos(dados: Path) -> Path:
    """O outubro sintético copia setembro com os mesmos ids. Na Orbe real, cada mês traz ids novos."""
    for arquivo in [*dados.glob("01_*.csv"), *dados.glob("06_*.csv")]:
        texto = arquivo.read_text(encoding="utf-8")
        arquivo.write_text(texto.replace("TX884", "TX984").replace("AN884", "AN984"), encoding="utf-8")
    return dados


def _com_atlas_na_planilha(dados: Path) -> None:
    """Comercial confirmou a Atlas: linha nova só com o CNPJ (sem EC), vigência a partir de 01/10."""
    caminho = next(dados.glob("02_*.xlsx"))
    wb = openpyxl.load_workbook(caminho)
    ws = wb["Taxas vigentes"]
    ultima = max(r for r in range(6, ws.max_row + 1) if ws.cell(r, 1).value and ws.cell(r, 2).value)
    ws.insert_rows(ultima + 1)
    valores = ["Atlas Idiomas Ltda", "39.566.044/0001-19", None, 0.0149, 0.0269, 0.0319, 0.0379, 0.0089,
               0.021, None, "a partir de 01/10/2026", None]
    for coluna, valor in enumerate(valores, start=1):
        ws.cell(ultima + 1, coluna).value = valor
    wb.save(caminho)


@pytest.fixture(scope="module")
def setembro_e_outubro(tmp_path_factory):
    base = tmp_path_factory.mktemp("dois_meses")
    saida = base / "saida"
    setembro = run_closing(DATA, saida, "2026-09")
    dados = _ids_novos(gerar_outubro(base / "dados"))
    _com_atlas_na_planilha(dados)
    return setembro, run_closing(dados, saida)


@pytest.fixture(scope="module")
def outubro_sem_contrato(tmp_path_factory):
    base = tmp_path_factory.mktemp("sem_contrato")
    saida = base / "saida"
    run_closing(DATA, saida, "2026-09")
    return run_closing(_ids_novos(gerar_outubro(base / "dados")), saida)


def _pend(r, tipo):
    return r.pendencias[r.pendencias["tipo"] == tipo]


def test_setembro_guarda_as_transacoes_sem_contrato_no_historico(setembro_e_outubro):
    setembro, _ = setembro_e_outubro
    guardadas = load_orbe(setembro.out_dir / NOME_PENDENTES)
    assert len(guardadas) == 57 and set(guardadas["cod_estabelecimento"]) == {ATLAS_EC}


def test_atlas_com_contrato_fatura_outubro_e_as_57_de_setembro(setembro_e_outubro):
    _, outubro = setembro_e_outubro
    assert outubro.problemas_integridade == [] and outubro.problemas_rastro.empty
    assert _pend(outubro, "SEM_CONTRATO").empty
    cm = outubro.rastro[outubro.rastro["id_saida"] == f"CK-CM-2026-10-{ATLAS_CNPJ}"]
    retro = cm[cm["arquivo_origem"] == ARQUIVO_PENDENTES]
    assert len(retro) == 57 and retro["metadata"].str.contains("retroativo_R18").all()
    assert (cm["arquivo_origem"] == "orbe").sum() > 0  # e as de outubro
    nota = outubro.notas.set_index("tomador_cnpj").loc[ATLAS_CNPJ, "valor_brl"]
    assert nota == sum(cm["valor"])  # setembro junto com outubro na mesma nota (R18)
    assert outubro.controle["orbe"]["fecha"]


def test_sem_contrato_as_de_setembro_seguem_na_pendencia(outubro_sem_contrato):
    pend = _pend(outubro_sem_contrato, "SEM_CONTRATO").iloc[0]
    assert "Inclui 57 transações de meses anteriores ainda pendentes" in pend["motivo"]
    rastro = outubro_sem_contrato.rastro
    assert (rastro[rastro["id_saida"] == pend["id_saida"]]["arquivo_origem"] == ARQUIVO_PENDENTES).sum() == 57
    guardadas = load_orbe(outubro_sem_contrato.out_dir / NOME_PENDENTES)
    assert len(guardadas) > 57  # setembro + outubro, para novembro
    assert outubro_sem_contrato.problemas_integridade == []


def test_alerta_usa_quem_estava_no_fechamento_anterior():
    txs = pd.DataFrame({"data_movimento": [pd.Timestamp("2026-10-05").date()]})  # 1ª venda depois do dia 1
    com_contrato = {"EC1"}
    assert _sem_historico("EC1", txs, "2026-10", com_contrato, "CNPJ1", {"CNPJ1"}) is False  # já estava
    assert _sem_historico("EC1", txs, "2026-10", com_contrato, "CNPJ1", {"OUTRO"}) is True  # novo
    assert _sem_historico("EC1", txs, "2026-10", com_contrato, "CNPJ1", None) is True  # sem histórico
    assert _sem_historico("EC9", txs, "2026-10", com_contrato, "CNPJ9", {"CNPJ9"}) is True  # sem contrato


def test_atlas_nova_no_mes_continua_com_alerta(setembro_e_outubro):
    _, outubro = setembro_e_outubro
    assert list(_pend(outubro, "ALERTA_ANTECIPACAO")["ref_origem"]) == [ATLAS_CNPJ]


def test_de_para_de_ec_vem_do_json(tmp_path):
    dados = json.loads(rules.CAMINHO_CADASTRO_EC.read_text(encoding="utf-8"))
    assert rules.EC_PARA_CNPJ == {ec: item["cnpj"] for ec, item in dados["ec_para_cnpj"].items()}
    assert set(rules.EC_SEM_CONTRATO) == set(dados["ec_sem_contrato"])
    novo = tmp_path / "cadastro.json"
    item = {"cnpj": "34671260000100", "fonte": "teste"}
    novo.write_text(json.dumps({"ec_para_cnpj": {"1099999999": item}}))
    assert rules.carregar_cadastro_ec(novo) == ({"1099999999": "34671260000100"}, {})
    assert rules.carregar_cadastro_ec(tmp_path / "nao_existe.json") == ({}, {})
