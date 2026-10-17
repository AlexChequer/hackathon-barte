"""Gera um pacote SINTÉTICO de outubro/2026 a partir dos arquivos reais de setembro.

Nada aqui é dado real: serve para testar o motor num mês que não temos (O1: "rodar em outubro
sem retrabalho"). Os arquivos de data/raw/ só são lidos, nunca alterados. Cenários inseridos:

- datas deslocadas 1 mês (setembro → outubro, agosto → setembro), nomes "out2026", sem o arquivo 06;
- lojista novo: 3 vendas de um EC que não está na planilha (EC_NOVO);
- anotação nova na planilha da Maple ("desconto acordado com o comercial até dezembro");
- repasse da Ponte referente a 09/2026 no Meridiano;
- ajuste de contestação da Ponte (linha nova, Spread R$ negativo, "AJUSTE REF PC-0926-0011", R27).
"""

from __future__ import annotations

import calendar
import datetime as dt
import email
import io
import re
from email import policy
from pathlib import Path

import openpyxl

ORIGEM = Path(__file__).resolve().parents[2] / "data" / "raw" / "set2026"
EC_NOVO = "1099999999"
EC_MODELO = "1041180753"  # Maple: as vendas do lojista novo copiam o formato das dela
CNPJ_MAPLE = "34671260000100"
AJUSTE_REF = "PC-0926-0011"  # Mundo Aberto: operação com spread acima do contrato em setembro
AJUSTE_SPREAD = -29.56
ANOTACAO_MAPLE = "desconto acordado com o comercial até dezembro"
DATA_BR = re.compile(r"\b(\d{2})/(\d{2})/(2026)\b")


def _mais_um_mes(d: dt.date) -> dt.date:
    ano, mes = (d.year + 1, 1) if d.month == 12 else (d.year, d.month + 1)
    return d.replace(year=ano, month=mes, day=min(d.day, calendar.monthrange(ano, mes)[1]))


def _texto_mais_um_mes(texto: str) -> str:
    def troca(m: re.Match) -> str:
        novo = _mais_um_mes(dt.date(int(m[3]), int(m[2]), int(m[1])))
        return f"{novo.day:02d}/{novo.month:02d}/{novo.year}"
    return DATA_BR.sub(troca, texto)


def _orbe(destino: Path) -> None:
    linhas = (ORIGEM / "01_orbe_relatorio_comissoes_set2026.csv").read_text(encoding="utf-8").splitlines()
    novas = [_texto_mais_um_mes(linha) for linha in linhas]
    modelo = [linha for linha in novas[1:] if f";{EC_MODELO};" in linha and ";VENDA;" in linha][:3]
    for i, linha in enumerate(modelo, start=1):
        campos = linha.split(";")
        campos[0], campos[2], campos[3] = f"TX9900000{i}", EC_NOVO, "ESCOLA NOVA IDIOMAS"
        novas.append(";".join(campos))
    arquivo = destino / "01_orbe_relatorio_comissoes_out2026.csv"
    arquivo.write_text("\r\n".join(novas) + "\r\n", encoding="utf-8")


def _taxas(destino: Path) -> None:
    wb = openpyxl.load_workbook(ORIGEM / "02_tabela_taxas_contratadas.xlsx")
    ws = wb["Taxas vigentes"]
    cabecalho = [c.value for c in ws[5]]
    col_cnpj, col_obs = cabecalho.index("CNPJ") + 1, cabecalho.index("Obs") + 1
    for linha in range(6, ws.max_row + 1):
        digitos = "".join(ch for ch in str(ws.cell(linha, col_cnpj).value or "") if ch.isdigit())
        if digitos.zfill(14) == CNPJ_MAPLE:
            ws.cell(linha, col_obs).value = ANOTACAO_MAPLE
    wb.save(destino / "02_tabela_taxas_contratadas.xlsx")


def _ponte(destino: Path) -> None:
    msg = email.message_from_bytes((ORIGEM / "03_email_ponte_cambio_relatorio_set2026.eml").read_bytes(),
                                   policy=policy.default)
    anexo = next(msg.iter_attachments())
    wb = openpyxl.load_workbook(io.BytesIO(anexo.get_content()))
    ws = wb["Operações"]
    linha_ref = None
    for row in ws.iter_rows():
        for c in row:
            if isinstance(c.value, dt.datetime):
                c.value = dt.datetime.combine(_mais_um_mes(c.value.date()), dt.time())
            elif isinstance(c.value, str):
                c.value = c.value.replace("09/2026", "10/2026").replace("PC-0926-", "PC-1026-")
        if row[0].value == AJUSTE_REF.replace("0926", "1026"):
            linha_ref = row
    rodape = next(r for r in range(6, ws.max_row + 1) if str(ws.cell(r, 1).value or "").startswith("Total"))
    ws.insert_rows(rodape - 1)  # antes da linha em branco que separa o rodapé
    for col, celula in enumerate(linha_ref, start=1):
        ws.cell(rodape - 1, col).value = celula.value
    ws.cell(rodape - 1, 1).value = "PC-1026-9001"
    ws.cell(rodape - 1, 11).value = AJUSTE_SPREAD  # Spread R$
    ws.cell(rodape - 1, 12).value = None  # Spread %
    ws.cell(rodape - 1, 16).value = f"AJUSTE REF {AJUSTE_REF}"
    ops = [r for r in ws.iter_rows(min_row=6, values_only=True) if r[0] and str(r[0]).startswith("PC-")]
    ws.cell(rodape + 1, 2).value = len(ops)
    ws.cell(rodape + 2, 2).value = round(sum(r[10] for r in ops if r[10] is not None), 2)
    buf = io.BytesIO()
    wb.save(buf)
    novo = email.message.EmailMessage()
    for chave in ("From", "To", "Date"):
        if msg[chave]:
            novo[chave] = msg[chave]
    novo["Subject"] = str(msg["Subject"]).replace("09/2026", "10/2026")
    novo.set_content(msg.get_body(("plain",)).get_content().replace("09/2026", "10/2026")
                     .replace("06/10/2026", "06/11/2026"))
    novo.add_attachment(buf.getvalue(), maintype="application",
                        subtype="vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                        filename="PonteCambio_Relatorio_TARVO_102026.xlsx")
    (destino / "03_email_ponte_cambio_relatorio_out2026.eml").write_bytes(novo.as_bytes())


def _aurora(destino: Path) -> None:
    texto = (ORIGEM / "05a_extrato_banco_aurora_cc4471_set2026.csv").read_text(encoding="utf-8")
    arquivo = destino / "05a_extrato_banco_aurora_cc4471_out2026.csv"
    arquivo.write_text(_texto_mais_um_mes(texto), encoding="utf-8")


def _meridiano(destino: Path) -> None:
    texto = (ORIGEM / "05b_extrato_banco_meridiano_cc90233_set2026.ofx").read_text(encoding="cp1252")
    texto = re.sub(r"(<DT\w+>)202609", r"\g<1>202610", texto)
    texto = texto.replace("REPASSE COMISSAO 08/2026", "REPASSE COMISSAO 09/2026")
    (destino / "05b_extrato_banco_meridiano_cc90233_out2026.ofx").write_text(texto, encoding="cp1252")


def gerar_outubro(destino: Path) -> Path:
    """Escreve o pacote sintético em `destino` e devolve a pasta."""
    destino.mkdir(parents=True, exist_ok=True)
    _orbe(destino)
    _taxas(destino)
    _ponte(destino)
    plano = "04_plano_de_contas_fenix_erp.csv"
    (destino / plano).write_bytes((ORIGEM / plano).read_bytes())
    _aurora(destino)
    _meridiano(destino)
    aviso = "Pacote SINTÉTICO de outubro/2026 gerado a partir de setembro (tests/sinteticos/outubro.py).\n"
    (destino / "00_LEIA-ME.md").write_text(aviso + "Não é dado real.\n", encoding="utf-8")
    return destino
