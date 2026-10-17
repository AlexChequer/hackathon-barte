"""Leitura do data pack como ele chega: relatório da Orbe (CSV) e planilha de taxas (xlsx).

Nenhuma regra de negócio aqui: só normalização de formato (número pt-BR, data, CNPJ, taxa).
"""

from __future__ import annotations

import csv
import re
from dataclasses import dataclass, field
from datetime import date, datetime
from decimal import Decimal
from pathlib import Path

import openpyxl

PADRAO_ORBE = "01_orbe_relatorio_comissoes_*.csv"
PADRAO_TAXAS = "02_tabela_taxas_contratadas*.xlsx"
ABA_TAXAS = "Taxas vigentes"

# Coluna da planilha -> chave interna da taxa.
COLUNAS_TAXA = {
    "Débito": "DEBITO",
    "Crédito à vista": "CREDITO_VISTA",
    "Créd. 2-6x": "CREDITO_2_6",
    "Créd. 7-12x": "CREDITO_7_12",
    "Pix": "PIX",
    "Antecipação": "ANTECIPACAO",
}


def achar_arquivo(pasta: Path, padrao: str) -> Path:
    achados = sorted(pasta.glob(padrao))
    if len(achados) != 1:
        raise FileNotFoundError(f"esperava 1 arquivo '{padrao}' em {pasta}, achei {len(achados)}")
    return achados[0]


def numero_br(texto: str) -> Decimal:
    """'-2.193,16' -> Decimal('-2193.16')."""
    return Decimal(texto.strip().replace(".", "").replace(",", "."))


def data_br(texto: str) -> date | None:
    texto = (texto or "").strip()
    return datetime.strptime(texto, "%d/%m/%Y").date() if texto else None


def normalizar_cnpj(valor) -> str | None:
    """Aceita '22.224.579/0001-09', '22224579000109' ou inteiro sem zero à esquerda."""
    if valor is None:
        return None
    digitos = re.sub(r"\D", "", str(valor))
    if not digitos or len(digitos) > 14:
        return None
    return digitos.zfill(14)


def taxa_planilha(valor) -> Decimal | None:
    """0.0149 (número) ou '1,49%' (texto) -> Decimal('0.0149'). Vazio ou texto inválido -> None."""
    if valor is None or valor == "":
        return None
    if isinstance(valor, (int, float)):
        return Decimal(repr(valor))
    texto = str(valor).strip()
    m = re.fullmatch(r"(-?\d+(?:,\d+)?)\s*%", texto)
    if not m:
        return None
    return Decimal(m.group(1).replace(",", ".")) / 100


def vigencia(valor) -> tuple[date | None, date | None] | None:
    """Retorna (início, fim). None quando a linha não tem vigência (ex.: lojista inativo)."""
    if valor is None or valor == "":
        return None
    if isinstance(valor, datetime):
        return valor.date(), None
    texto = str(valor).strip().lower()
    m = re.search(r"(\d{2}/\d{2}/\d{4})", texto)
    if not m:
        return None
    d = data_br(m.group(1))
    if texto.startswith("até"):
        return None, d
    return d, None


@dataclass
class LinhaOrbe:
    linha_arquivo: int
    id_transacao: str
    data_movimento: date
    ec: str
    nome_orbe: str
    tipo: str
    modalidade: str
    parcelas: int | None
    valor_bruto: Decimal
    taxa_aplicada_pct: Decimal
    valor_taxa: Decimal
    valor_liquido: Decimal
    id_referencia: str
    data_venda_original: date | None


@dataclass
class LinhaTaxa:
    linha_planilha: int
    lojista: str
    cnpj: str | None
    ec: str | None
    taxas: dict[str, Decimal | None]
    vigencia: tuple[date | None, date | None] | None
    obs: str
    problemas: list[str] = field(default_factory=list)


def ler_orbe(caminho: Path) -> list[LinhaOrbe]:
    linhas = []
    with caminho.open(encoding="utf-8-sig", newline="") as f:
        for n, r in enumerate(csv.DictReader(f, delimiter=";"), start=2):
            linhas.append(LinhaOrbe(
                linha_arquivo=n,
                id_transacao=r["id_transacao"].strip(),
                data_movimento=data_br(r["data_movimento"]),
                ec=r["cod_estabelecimento"].strip(),
                nome_orbe=r["nome_estabelecimento"].strip(),
                tipo=r["tipo_lancamento"].strip(),
                modalidade=r["modalidade"].strip(),
                parcelas=int(r["qtd_parcelas"]) if r["qtd_parcelas"].strip() else None,
                valor_bruto=numero_br(r["valor_bruto"]),
                taxa_aplicada_pct=numero_br(r["taxa_aplicada_pct"]),
                valor_taxa=numero_br(r["valor_taxa"]),
                valor_liquido=numero_br(r["valor_liquido"]),
                id_referencia=r["id_referencia"].strip(),
                data_venda_original=data_br(r["data_venda_original"]),
            ))
    return linhas


def ler_taxas(caminho: Path) -> list[LinhaTaxa]:
    """Lê a aba vigente. O cabeçalho é a linha cuja 1ª célula é 'Lojista'; os dados vão até a 1ª linha vazia."""
    wb = openpyxl.load_workbook(caminho, data_only=True)
    ws = wb[ABA_TAXAS]
    cab_linha, cabecalho = None, None
    for row in ws.iter_rows():
        if row[0].value and str(row[0].value).strip() == "Lojista":
            cab_linha = row[0].row
            cabecalho = [str(c.value).strip() if c.value is not None else "" for c in row]
            break
    if cab_linha is None:
        raise ValueError(f"cabeçalho 'Lojista' não encontrado na aba '{ABA_TAXAS}'")
    col = {nome: i for i, nome in enumerate(cabecalho) if nome}
    for obrigatoria in ("Lojista", "CNPJ", "EC Orbe", "Vigência", *COLUNAS_TAXA):
        if obrigatoria not in col:
            raise ValueError(f"coluna '{obrigatoria}' ausente na aba '{ABA_TAXAS}'")

    linhas = []
    for row in ws.iter_rows(min_row=cab_linha + 1, values_only=True):
        lojista = row[col["Lojista"]]
        if lojista is None or str(lojista).strip() == "":
            break
        problemas = []
        cnpj = normalizar_cnpj(row[col["CNPJ"]])
        if cnpj is None:
            problemas.append(f"CNPJ ilegível: {row[col['CNPJ']]!r}")
        ec_bruto = row[col["EC Orbe"]]
        ec = re.sub(r"\D", "", str(ec_bruto)) if ec_bruto is not None else ""
        taxas = {}
        for nome, chave in COLUNAS_TAXA.items():
            bruto = row[col[nome]]
            taxas[chave] = taxa_planilha(bruto)
            if bruto not in (None, "") and taxas[chave] is None:
                problemas.append(f"taxa '{nome}' ilegível: {bruto!r}")
        linhas.append(LinhaTaxa(
            linha_planilha=len(linhas) + cab_linha + 1,
            lojista=str(lojista).strip(),
            cnpj=cnpj,
            ec=ec or None,
            taxas=taxas,
            vigencia=vigencia(row[col["Vigência"]]),
            obs=str(row[col["Obs"]]).strip() if "Obs" in col and row[col["Obs"]] else "",
            problemas=problemas,
        ))
    return linhas
